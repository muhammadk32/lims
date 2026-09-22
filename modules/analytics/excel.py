"""Excel export helpers for analytics reports."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


_BOLD_WHITE = Font(bold=True, color='FFFFFF')
_BOLD = Font(bold=True)
_HEADER_FILL = PatternFill('solid', fgColor='212529')
_KPI_FILL = PatternFill('solid', fgColor='F8F9FA')
_TOTAL_FILL = PatternFill('solid', fgColor='E9ECEF')
_CENTER = Alignment(horizontal='center', vertical='center')
_RIGHT = Alignment(horizontal='right')
_THIN = Side(style='thin', color='6C757D')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def build_workbook(title, subtitle, kpis, headers, rows, widths=None,
                   number_cols=None, total_row=None):
    """Build an Excel workbook. Returns a BytesIO buffer ready for send_file.

    kpis         : list of (label, value, number_format) or None
    headers      : list of column titles
    rows         : list of lists
    widths       : list of column widths
    number_cols  : 1-based column indices that are numeric
    total_row    : optional list matching headers (or None for blanks)
    """
    number_cols = number_cols or []
    wb = Workbook()
    ws = wb.active
    ws.title = (title or 'Report')[:30]

    n_cols = len(headers)
    last_col = get_column_letter(n_cols)

    # Title
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.merge_cells(f'A1:{last_col}1')
    ws.cell(row=2, column=1, value=subtitle).font = Font(size=10, italic=True, color='6C757D')
    ws.merge_cells(f'A2:{last_col}2')

    row_ptr = 4

    # KPI block
    if kpis:
        for label, val, fmt in kpis:
            a = ws.cell(row=row_ptr, column=1, value=label)
            a.font = _BOLD
            a.fill = _KPI_FILL
            a.border = _BORDER
            b = ws.cell(row=row_ptr, column=2, value=val)
            if fmt:
                b.number_format = fmt
            b.fill = _KPI_FILL
            b.border = _BORDER
            row_ptr += 1
        row_ptr += 1

    # Header row
    header_row = row_ptr
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = _BOLD_WHITE
        c.fill = _HEADER_FILL
        c.alignment = _CENTER
        c.border = _BORDER
    row_ptr += 1

    # Data rows
    for r in rows:
        for col, v in enumerate(r, 1):
            c = ws.cell(row=row_ptr, column=col, value=v)
            c.border = _BORDER
            if col in number_cols:
                c.number_format = '#,##0.00'
                c.alignment = _RIGHT
        row_ptr += 1

    # Total row
    if total_row:
        for col, v in enumerate(total_row, 1):
            c = ws.cell(row=row_ptr, column=col, value=v)
            c.font = _BOLD
            c.fill = _TOTAL_FILL
            c.border = _BORDER
            if col in number_cols:
                c.number_format = '#,##0.00'
                c.alignment = _RIGHT
        row_ptr += 1

    # Column widths
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
